"""Unit tests for Word, PowerPoint, and Paint sketch desktop automation tools."""

import os
from pathlib import Path
import pytest
from unittest.mock import patch, MagicMock

import docx
import pptx

from backend.tools.office.word import (
    write_word_document,
    generate_word_content,
    _parse_generated_word_text,
    _fallback_word_content,
)
from backend.tools.office.powerpoint import (
    create_powerpoint_presentation,
    generate_presentation_content,
    _parse_generated_pptx_text,
    _fallback_presentation_content,
)
from backend.tools.desktop.paint import (
    draw_sketch_in_paint,
    draw_house_sketch,
    draw_tree_sketch,
    draw_sun_sketch,
    draw_landscape_sketch,
    draw_shape_sketch,
)


def test_word_content_parsing():
    raw_text = """TITLE: Artificial Intelligence Overview
SUBTITLE: A Comprehensive Guide to Modern AI

HEADING: 1. Introduction
CONTENT: Artificial Intelligence is rapidly reshaping technology.

HEADING: 2. Machine Learning Basics
CONTENT: Machine learning algorithms learn patterns from data.
"""
    parsed = _parse_generated_word_text(raw_text, "Artificial Intelligence")
    assert parsed["title"] == "Artificial Intelligence Overview"
    assert parsed["subtitle"] == "A Comprehensive Guide to Modern AI"
    assert len(parsed["sections"]) == 2
    assert parsed["sections"][0]["heading"] == "1. Introduction"


def test_word_fallback_content():
    fb = _fallback_word_content("Quantum Computing", 500)
    assert fb["title"] == "Quantum Computing"
    assert len(fb["sections"]) >= 4


@patch("backend.tools.office.word.subprocess.Popen")
@patch("backend.tools.office.word.os.startfile", create=True)
def test_write_word_document_execution(mock_startfile, mock_popen, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    res = write_word_document.invoke({
        "topic": "Artificial Intelligence",
        "approx_word_count": 300,
        "file_name": "ai_test.docx",
    })
    assert "Successfully generated and wrote topic 'Artificial Intelligence'" in res
    assert "Saved File Path:" in res

    # Verify document exists on disk and has content
    file_path = tmp_path / "data" / "output" / "ai_test.docx"
    assert file_path.exists()
    doc = docx.Document(str(file_path))
    assert len(doc.paragraphs) > 0


def test_pptx_content_parsing():
    raw_text = """PRESENTATION_TITLE: Machine Learning Intro
PRESENTATION_SUBTITLE: Key Concepts & Real World Apps

SLIDE: Introduction to ML
BULLET: Supervised Learning
BULLET: Unsupervised Learning

SLIDE: Deep Learning
BULLET: Neural Networks
BULLET: Convolutional Nets
"""
    parsed = _parse_generated_pptx_text(raw_text, "Machine Learning", num_slides=3)
    assert parsed["title"] == "Machine Learning Intro"
    assert len(parsed["slides"]) == 2
    assert parsed["slides"][0]["title"] == "Introduction to ML"
    assert len(parsed["slides"][0]["bullets"]) == 2


def test_pptx_fallback_content():
    fb = _fallback_presentation_content("Robotics", num_slides=5)
    assert fb["title"] == "Robotics"
    assert len(fb["slides"]) == 4


@patch("backend.tools.office.powerpoint.subprocess.Popen")
@patch("backend.tools.office.powerpoint.os.startfile", create=True)
def test_create_powerpoint_presentation_execution(mock_startfile, mock_popen, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    res = create_powerpoint_presentation.invoke({
        "topic": "Artificial Intelligence",
        "num_slides": 5,
        "file_name": "ai_deck.pptx",
    })
    assert "Successfully created a 5-slide PowerPoint presentation" in res
    assert "ai_deck.pptx" in res

    # Verify presentation file exists on disk and can be read by pptx
    file_path = tmp_path / "data" / "output" / "ai_deck.pptx"
    assert file_path.exists()
    prs = pptx.Presentation(str(file_path))
    assert len(prs.slides) == 5


def test_paint_sketch_routines():
    house_acts = draw_house_sketch()
    assert len(house_acts) >= 3

    sun_acts = draw_sun_sketch()
    assert len(sun_acts) >= 2

    landscape_acts = draw_landscape_sketch()
    assert len(landscape_acts) >= 4

    shape_acts = draw_shape_sketch("circle")
    assert len(shape_acts) >= 1


@patch("backend.tools.desktop.paint.run_system_app")
@patch("backend.tools.desktop.paint.window_controller")
@patch("backend.tools.desktop.paint.screen_capture")
@patch("backend.tools.desktop.paint.mouse_controller")
def test_draw_sketch_in_paint_tool(mock_mouse, mock_screen, mock_window, mock_run_app):
    mock_run_app.return_value = "Successfully launched system application 'paint'"
    mock_window.get_active_window.return_value = {"title": "Paint", "width": 1920, "height": 1080}
    mock_screen.capture_frame.return_value = {"data": "base64_data", "width": 800, "height": 600}

    res = draw_sketch_in_paint.invoke({"sketch_type": "house"})
    assert "Successfully executed sketch drawing in Microsoft Paint" in res
    assert "Sketch Requested: 'house'" in res
    assert "PASSED" in res
